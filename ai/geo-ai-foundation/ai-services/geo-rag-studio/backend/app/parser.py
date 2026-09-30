import asyncio
import csv
import io
import json
import mimetypes
import tempfile
import time
import uuid
import zipfile
from pathlib import Path
from urllib.parse import urlsplit

import fitz
import httpx

from .config import settings
from .db import new_id


def parse_pdf_layout(path: str) -> list[dict]:
    """Extract selectable PDF text and its page-space bounding boxes."""
    blocks = []
    with fitz.open(path) as document:
        ordinal = 0
        for page_number, page in enumerate(document, start=1):
            rect = page.rect
            for raw in page.get_text("blocks", sort=True):
                x0, y0, x1, y1, text, *_ = raw
                text = (text or "").strip()
                if not text:
                    continue
                blocks.append(
                    {
                        "id": new_id("BLK"),
                        "ordinal": ordinal,
                        "page": page_number,
                        "text": text,
                        "bbox": [float(x0), float(y0), float(x1), float(y1)],
                        "page_width": float(rect.width),
                        "page_height": float(rect.height),
                    }
                )
                ordinal += 1
    return blocks


def parse_text(text: str) -> list[dict]:
    return [
        {
            "id": new_id("BLK"),
            "ordinal": index,
            "page": None,
            "text": part,
            "bbox": None,
            "page_width": None,
            "page_height": None,
        }
        for index, part in enumerate(part.strip() for part in text.split("\n\n") if part.strip())
    ]


def _markdown_table(rows: list[list[object]]) -> str:
    normalized = [[str(value or "").replace("|", "\\|").replace("\n", " ").strip() for value in row] for row in rows]
    while normalized and not any(normalized[-1]):
        normalized.pop()
    if not normalized:
        return ""
    width = max(map(len, normalized))
    normalized = [row + [""] * (width - len(row)) for row in normalized]
    return "\n".join([f"| {' | '.join(normalized[0])} |", f"| {' | '.join(['---'] * width)} |", *[f"| {' | '.join(row)} |" for row in normalized[1:]]])


def _parse_csv(path: Path) -> str:
    try:
        text = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        text = path.read_text(encoding="gb18030")
    rows = list(csv.reader(io.StringIO(text)))
    return _markdown_table(rows)


def _store_embedded_image(image_data: bytes, extension: str, knowledge_base_id: str | None) -> str | None:
    if not knowledge_base_id:
        return None
    from .config import settings

    kb_id = "".join(char for char in knowledge_base_id if char.isalnum() or char in "-_")
    if not kb_id:
        return None
    object_id = uuid.uuid4().hex
    suffix = extension.lower().lstrip(".") or "bin"
    destination = Path(settings.rag_upload_dir) / ".kb-images" / kb_id / f"{object_id}.{suffix}"
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_bytes(image_data)
    return f"/api/knowledge/databases/{kb_id}/images/kb-images/{object_id}.{suffix}"


def _parse_docx(path: Path, knowledge_base_id: str | None = None) -> str:
    from docx import Document

    document = Document(str(path))
    blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        rows = [[cell.text.replace("\n", " ").strip() for cell in row.cells] for row in table.rows]
        table_markdown = _markdown_table(rows)
        if table_markdown:
            blocks.append(table_markdown)
    for image_part in document.part.package.image_parts:
        url = _store_embedded_image(image_part.blob, image_part.filename.rsplit(".", 1)[-1], knowledge_base_id)
        if url:
            blocks.append(f"![{Path(image_part.filename).name}]({url})")
    return "\n\n".join(blocks)


def _parse_pptx(path: Path, knowledge_base_id: str | None = None) -> str:
    from pptx import Presentation

    presentation = Presentation(str(path))
    slides = []
    for index, slide in enumerate(presentation.slides, start=1):
        blocks = [f"## 第 {index} 页"]
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = "\n".join(paragraph.text for paragraph in shape.text_frame.paragraphs).strip()
                if text:
                    blocks.append(text)
            if getattr(shape, "has_table", False):
                table = _markdown_table([[cell.text for cell in row.cells] for row in shape.table.rows])
                if table:
                    blocks.append(table)
            if getattr(shape, "shape_type", None) == 13:
                image = shape.image
                url = _store_embedded_image(image.blob, image.ext, knowledge_base_id)
                if url:
                    blocks.append(f"![{Path(image.filename).name}]({url})")
        slides.append("\n\n".join(blocks))
    return "\n\n".join(slides)


def _parse_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = []
        for worksheet in workbook.worksheets:
            rows = [[value for value in row] for row in worksheet.iter_rows(values_only=True)]
            table = _markdown_table(rows)
            if table:
                sheets.append(f"## {worksheet.title}\n\n{table}")
        return "\n\n".join(sheets)
    finally:
        workbook.close()


def _parse_xls(path: Path) -> str:
    import xlrd

    workbook = xlrd.open_workbook(path)
    sheets = []
    for worksheet in workbook.sheets():
        table = _markdown_table([worksheet.row_values(index) for index in range(worksheet.nrows)])
        if table:
            sheets.append(f"## {worksheet.name}\n\n{table}")
    return "\n\n".join(sheets)


def _parse_local_document(path: Path, parser_options: dict | None = None) -> tuple[list[dict], str]:
    parser_options = parser_options or {}
    knowledge_base_id = parser_options.get("knowledge_base_id")
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        try:
            text = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            text = path.read_text(encoding="gb18030")
        return parse_text(text), "plain-text"
    if suffix in {".html", ".htm"}:
        from markdownify import markdownify

        text = markdownify(path.read_text(encoding="utf-8", errors="replace"), heading_style="ATX")
    elif suffix == ".json":
        value = json.loads(path.read_text(encoding="utf-8-sig"))
        text = "```json\n" + json.dumps(value, ensure_ascii=False, indent=2) + "\n```"
    elif suffix == ".csv":
        text = _parse_csv(path)
    elif suffix == ".docx":
        text = _parse_docx(path, knowledge_base_id)
    elif suffix == ".pptx":
        text = _parse_pptx(path, knowledge_base_id)
    elif suffix == ".xlsx":
        text = _parse_xlsx(path)
    elif suffix == ".xls":
        text = _parse_xls(path)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")
    return parse_text(text), "local-structured-parser"


async def parse_file(path: str, parser_options: dict | None = None) -> tuple[list[dict], str]:
    """Parse Yuxi-compatible local document formats without claiming OCR when no engine is configured."""
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    options = parser_options or {}
    engine = str(options.get("engine") or "auto").strip().lower()
    api_uri = options.get("mineru_api_uri") or None
    api_key = options.get("api_key") or None
    if suffix == ".pdf":
        if engine == "mineru_official" or (engine == "auto" and (api_key or settings.mineru_api_key)):
            blocks = await parse_pdf_mineru_official(str(file_path), api_key, api_uri)
            return blocks, "mineru-official"
        if engine == "mineru" or (engine == "auto" and (api_uri or (settings.mineru_enabled and settings.mineru_api_uri))):
            blocks = await parse_pdf_mineru(str(file_path), api_uri)
            return blocks, "mineru-file_parse"
        return parse_pdf_layout(str(file_path)), "pymupdf-layout"
    if suffix in {".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".tif", ".webp"}:
        selected_engine = "mineru_official" if engine == "mineru_official" or (engine == "auto" and (api_key or settings.mineru_api_key)) else "mineru" if engine == "mineru" or (engine == "auto" and (api_uri or (settings.mineru_enabled and settings.mineru_api_uri))) else ""
        if not selected_engine:
            raise ValueError("图片文件需要可用的 OCR 服务，请配置 MinerU API 或选择已配置的 OCR 引擎")
        parse_path = file_path
        temp_path = None
        if suffix not in {".png", ".jpg", ".jpeg"}:
            from PIL import Image

            Path(settings.rag_upload_dir).mkdir(parents=True, exist_ok=True)
            with tempfile.NamedTemporaryFile(delete=False, suffix=".png", dir=settings.rag_upload_dir) as temp_file:
                temp_path = Path(temp_file.name)
            with Image.open(file_path) as image:
                image.convert("RGB").save(temp_path, format="PNG")
            parse_path = temp_path
        try:
            if selected_engine == "mineru_official":
                blocks = await parse_pdf_mineru_official(str(parse_path), api_key, api_uri)
            else:
                blocks = await parse_pdf_mineru(str(parse_path), api_uri)
            return blocks, "mineru-official" if selected_engine == "mineru_official" else "mineru-file_parse"
        finally:
            if temp_path:
                temp_path.unlink(missing_ok=True)
    if suffix in {".docx", ".pptx"} and engine == "mineru_official":
        return await parse_pdf_mineru_official(str(file_path), api_key, api_uri), "mineru-official"
    return await asyncio.to_thread(_parse_local_document, file_path, options)


async def parse_file_markdown(path: str, parser_options: dict | None = None) -> str:
    blocks, _ = await parse_file(path, parser_options)
    return "\n\n".join(block["text"] for block in blocks)


def _mineru_blocks(payload: bytes, pdf_path: str) -> list[dict]:
    page_sizes = {}
    try:
        with fitz.open(pdf_path) as document:
            page_sizes = {
                index: (float(page.rect.width), float(page.rect.height))
                for index, page in enumerate(document)
            }
    except (RuntimeError, ValueError):
        # MinerU can parse office files too; those have page references but no local PDF geometry.
        page_sizes = {}

    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        names = archive.namelist()
        content_list_name = next((name for name in names if name.endswith("content_list.json")), None)
        if content_list_name:
            content = json.loads(archive.read(content_list_name).decode("utf-8"))
            items = content if isinstance(content, list) else content.get("content_list", [])
            blocks = []
            for ordinal, item in enumerate(items):
                if not isinstance(item, dict):
                    continue
                text = str(item.get("text") or item.get("content") or "").strip()
                if not text:
                    continue
                page_index = item.get("page_idx", item.get("page"))
                if isinstance(page_index, str) and page_index.isdigit():
                    page_index = int(page_index)
                if type(page_index) is not int:
                    page_index = 0
                page_number = page_index + 1 if "page_idx" in item else max(page_index, 1)
                width, height = page_sizes.get(page_number - 1, (None, None))
                bbox = item.get("bbox")
                if not (isinstance(bbox, list) and len(bbox) == 4 and width and height):
                    bbox = None
                else:
                    try:
                        coords = [float(value) for value in bbox]
                        x0, y0, x1, y1 = coords
                        if x0 < 0 or y0 < 0 or x1 <= x0 or y1 <= y0 or x1 > width + 1 or y1 > height + 1:
                            coords = None
                    except (TypeError, ValueError):
                        coords = None
                    # Unknown MinerU coordinate scale: keep the text/page but don't draw a false locator.
                    bbox = coords
                blocks.append(
                    {
                        "id": new_id("BLK"),
                        "ordinal": ordinal,
                        "page": page_number,
                        "text": text,
                        "bbox": bbox,
                        "page_width": width,
                        "page_height": height,
                    }
                )
            if blocks:
                return blocks

        markdown_name = next((name for name in names if name.endswith(".md")), None)
        if markdown_name:
            return parse_text(archive.read(markdown_name).decode("utf-8", errors="replace"))
    raise ValueError("MinerU response did not contain extractable text")


async def parse_pdf_mineru(path: str, api_uri: str | None = None) -> list[dict]:
    endpoint_uri = api_uri or (settings.mineru_api_uri if settings.mineru_enabled else "")
    if not endpoint_uri:
        raise ValueError("MinerU OCR is not configured; set the knowledge-base MinerU API URL or MINERU_API_URI")
    endpoint = f"{endpoint_uri.rstrip('/')}/file_parse"
    data = {
        "lang_list": ["ch"],
        "backend": "hybrid-auto-engine",
        "parse_method": "auto",
        "return_md": "true",
        "response_format_zip": "true",
        "return_images": "true",
    }
    async with httpx.AsyncClient(timeout=settings.mineru_timeout_seconds) as client:
        with Path(path).open("rb") as source:
            response = await client.post(
                endpoint,
                data=data,
                files={"files": (Path(path).name, source, mimetypes.guess_type(path)[0] or "application/octet-stream")},
            )
    response.raise_for_status()
    return _mineru_blocks(response.content, path)


async def parse_pdf_mineru_official(
    path: str,
    api_key: str | None = None,
    api_base: str | None = None,
) -> list[dict]:
    """Upload a PDF to MinerU's official batch API and parse its returned ZIP."""
    key = (api_key or settings.mineru_api_key).strip()
    if not key:
        raise ValueError("MinerU 官方 API 未配置 API Key")
    base = (api_base or "https://mineru.net/api/v4").strip().rstrip("/")
    parsed = urlsplit(base)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise ValueError("MinerU API Base URL 必须是有效的 http 或 https 地址")
    if parsed.scheme != "https" and parsed.hostname not in {"localhost", "127.0.0.1", "::1"}:
        raise ValueError("向非本机 MinerU 服务发送 API Key 时必须使用 HTTPS")

    file_path = Path(path)
    file_name = file_path.name
    data_id = f"geo-rag-{uuid.uuid4().hex}"
    headers = {"Authorization": f"Bearer {key}"}
    async with httpx.AsyncClient(timeout=settings.mineru_timeout_seconds) as client:
        response = await client.post(
            f"{base}/file-urls/batch",
            headers=headers,
            json={
                "enable_formula": True,
                "enable_table": True,
                "language": "ch",
                "files": [{"name": file_name, "is_ocr": True, "data_id": data_id}],
            },
        )
        response.raise_for_status()
        payload = response.json()
        if payload.get("code") != 0:
            raise ValueError(f"MinerU 申请上传链接失败：{payload.get('msg') or '未知错误'}")
        data = payload.get("data") or {}
        batch_id = data.get("batch_id")
        upload_urls = data.get("file_urls") or []
        if not batch_id or not upload_urls:
            raise ValueError("MinerU 未返回 batch_id 或上传链接")

        upload_response = await client.put(
            upload_urls[0],
            content=file_path.read_bytes(),
            timeout=60.0,
        )
        upload_response.raise_for_status()

        deadline = time.monotonic() + settings.mineru_timeout_seconds
        file_result = None
        while time.monotonic() < deadline:
            result_response = await client.get(
                f"{base}/extract-results/batch/{batch_id}",
                headers=headers,
                timeout=30.0,
            )
            result_response.raise_for_status()
            result_payload = result_response.json()
            if result_payload.get("code") != 0:
                raise ValueError(f"MinerU 查询解析状态失败：{result_payload.get('msg') or '未知错误'}")
            results = (result_payload.get("data") or {}).get("extract_result") or []
            file_result = next((item for item in results if item.get("data_id") == data_id), None)
            if file_result is None and results:
                file_result = results[0]
            state = (file_result or {}).get("state")
            if state == "done":
                break
            if state == "failed":
                raise ValueError(f"MinerU 文档解析失败：{file_result.get('err_msg') or '未知错误'}")
            await asyncio.sleep(2)
        else:
            raise ValueError("MinerU 文档解析超时")

        zip_url = (file_result or {}).get("full_zip_url")
        if not zip_url:
            raise ValueError("MinerU 未返回解析结果 ZIP 地址")
        zip_response = await client.get(zip_url, timeout=settings.mineru_timeout_seconds)
        zip_response.raise_for_status()
        return _mineru_blocks(zip_response.content, str(file_path))


def render_pdf_page(path: str, page_no: int, zoom: float = 1.6) -> bytes:
    with fitz.open(path) as document:
        if page_no < 1 or page_no > document.page_count:
            raise ValueError("page out of range")
        page = document.load_page(page_no - 1)
        pixmap = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        return pixmap.tobytes("png")
