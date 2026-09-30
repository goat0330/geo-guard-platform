import io

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation

from app.config import settings
from app.main import app


def _client(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "rag_db_path", str(tmp_path / "rag.sqlite3"))
    monkeypatch.setattr(settings, "rag_upload_dir", str(tmp_path / "uploads"))
    monkeypatch.setattr(settings, "embedding_base_url", "")
    monkeypatch.setattr(settings, "embedding_api_key", "")
    monkeypatch.setattr(settings, "embedding_model", "")
    monkeypatch.setattr(settings, "rerank_base_url", "")
    monkeypatch.setattr(settings, "rerank_api_key", "")
    monkeypatch.setattr(settings, "rerank_model", "")
    monkeypatch.setattr(settings, "mineru_enabled", False)
    monkeypatch.setattr(settings, "mineru_api_uri", "")
    monkeypatch.setattr(settings, "mineru_api_key", "")
    return TestClient(app)


def _png_bytes():
    image = Image.new("RGB", (8, 8), color="red")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def test_office_csv_html_and_json_uploads_parse_and_index(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    kb = client.post("/api/v1/knowledge-bases", json={"name": "多格式验收", "kb_type": "local"}).json()
    kb_id = kb["id"]

    docx = Document()
    docx.add_paragraph("DOCX 灾害巡查正文")
    table = docx.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "记录"
    table.cell(1, 0).text = "坡体裂缝"
    table.cell(1, 1).text = "持续扩大"
    docx.add_picture(io.BytesIO(_png_bytes()))
    docx_buffer = io.BytesIO()
    docx.save(docx_buffer)

    pptx = Presentation()
    slide = pptx.slides.add_slide(pptx.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, 4000000, 500000).text = "PPTX 应急避险流程"
    pptx_buffer = io.BytesIO()
    pptx.save(pptx_buffer)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "巡查记录"
    worksheet.append(["隐患点", "现象"])
    worksheet.append(["北侧边坡", "雨后渗水"])
    xlsx_buffer = io.BytesIO()
    workbook.save(xlsx_buffer)

    samples = [
        ("巡查.docx", docx_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "坡体裂缝"),
        ("预案.pptx", pptx_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation", "应急避险流程"),
        ("台账.xlsx", xlsx_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "雨后渗水"),
        ("网页.html", "<h1>网页巡查</h1><p>发现落石风险</p>".encode("utf-8"), "text/html", "落石风险"),
        ("记录.csv", "点位,现象\n南侧坡脚,裂缝扩展\n".encode("utf-8"), "text/csv", "裂缝扩展"),
        ("规则.json", '{"风险":"高","依据":"持续变形"}'.encode("utf-8"), "application/json", "持续变形"),
    ]
    indexed_ids = []
    for filename, content, mime_type, expected_text in samples:
        response = client.post(
            "/api/v1/documents/upload",
            data={"knowledge_base_id": kb_id},
            files={"file": (filename, content, mime_type)},
        )
        assert response.status_code == 200, response.text
        assert response.json()["status"] == "indexed"
        assert response.json()["chunks"] > 0
        indexed_ids.append(response.json()["document_id"])
        chunk_text = " ".join(chunk["text"] for chunk in client.get(f"/api/v1/documents/{response.json()['document_id']}/chunks").json())
        assert expected_text in chunk_text

    docx_content = client.get(f"/api/knowledge/databases/{kb_id}/documents/{indexed_ids[0]}/content").json()
    image_url = next(
        part.split(")", 1)[0]
        for block in docx_content["blocks"]
        for part in block["text"].split("(")[1:]
        if part.startswith(f"/api/knowledge/databases/{kb_id}/images/kb-images/")
    )
    image = client.get(image_url)
    assert image.status_code == 200
    assert image.headers["content-type"] == "image/png"
    assert image.content == _png_bytes()

    exported = client.get(f"/api/knowledge/databases/{kb_id}/export", params={"format": "csv"})
    assert exported.status_code == 200
    assert b"chunk_id" in exported.content
    workbook_export = client.get(f"/api/knowledge/databases/{kb_id}/export", params={"format": "xlsx"})
    assert workbook_export.status_code == 200
    assert workbook_export.content[:2] == b"PK"


def test_markdown_preview_and_image_ocr_unavailable_are_explicit(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    supported = client.get("/api/knowledge/files/supported-types")
    assert supported.status_code == 200
    assert {".docx", ".pptx", ".xls", ".xlsx", ".html", ".png", ".zip"}.issubset(supported.json()["file_types"])
    preview = client.post(
        "/api/knowledge/files/markdown",
        files={"file": ("预案.html", "<h2>避险流程</h2><p>立即撤离</p>".encode("utf-8"), "text/html")},
    )
    assert preview.status_code == 200, preview.text
    assert "避险流程" in preview.json()["markdown_content"]
    assert "立即撤离" in preview.json()["markdown_content"]

    kb = client.post("/api/v1/knowledge-bases", json={"name": "图片 OCR 失败记录", "kb_type": "local"}).json()
    failed = client.post(
        "/api/v1/documents/upload",
        data={"knowledge_base_id": kb["id"]},
        files={"file": ("现场.png", _png_bytes(), "image/png")},
    )
    assert failed.status_code == 422
    documents = client.get(f"/api/v1/knowledge-bases/{kb['id']}/documents").json()
    assert documents[0]["status"] == "failed"
    assert "OCR" in documents[0]["error_message"]
